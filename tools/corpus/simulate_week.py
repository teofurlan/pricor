"""Simulate "next week" for a real supplier price list, for demos and tests.

Usage (from the repository root)::

    uv run python -m tools.corpus.simulate_week data/raw/<id>/<date>/<file> \
        [--out-dir DIR] [--seed N] [--added N] [--removed N] [--move-column] [--anomaly]

Reads one real snapshot and writes a *clearly labeled* variant next to a JSON manifest that
lists every modification: price changes (with the percentage), added products, removed codes,
a moved column and a planted anomaly. Everything is derived from `--seed`, so the same inputs
always give the same simulation. The output is always `.xlsx` (a legacy `.xls` source is
converted) and says SIMULATED in its file name, in a `SIMULATED` sheet and in the workbook
properties; it is never real supplier data and must never be presented as such.

The output folder must be outside the repository or under `data/raw/` (gitignored). A snapshot
under `data/raw/<id>/<date>/` defaults to `data/raw/<id>/<date>-simulated/`.
"""

import argparse
import hashlib
import json
import random
import re
import sys
from dataclasses import asdict, dataclass, replace
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from pricor.adapters.heuristic_classifier import HeuristicColumnClassifier
from pricor.adapters.spreadsheet_reader import FileSpreadsheetReader
from pricor.domain.parsing import CellError, parse_number
from pricor.domain.profile import ColumnProfile, RowClass, profile_sheet
from pricor.domain.roles import ColumnRole

REPO_ROOT = Path(__file__).resolve().parents[2]
RAW_ROOT = REPO_ROOT / "data" / "raw"
TOOL = "tools.corpus.simulate_week"
MARKER = "SIMULATED"
SHEET_NAME_LIMIT = 31
_FIXED_STAMP = datetime(2026, 1, 1)  # keeps the workbook properties identical between runs
_NUMBER_TOKEN = re.compile(r"[0-9][0-9.,]*")
_BAD_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
_CENT = Decimal("0.01")


@dataclass(frozen=True)
class SimulationOptions:
    seed: int = 1
    drift_share: Decimal = Decimal("0.6")  # share of the priced products whose price moves
    drift_min: Decimal = Decimal("0.01")  # smallest drift, as a ratio (0.01 is 1%)
    drift_max: Decimal = Decimal("0.08")
    down_share: Decimal = Decimal("0.2")  # share of the moved prices that go down
    added: int = 0
    removed: int = 0
    move_column: bool = False  # move the price column one position
    anomaly: bool = False  # one product whose price jumps by `anomaly_factor`
    anomaly_factor: Decimal = Decimal("3")
    min_price: Decimal = Decimal("1")  # cheaper products stay as they are: cents would swallow the drift
    code_col: int | None = None  # 0-based; detected when not given
    price_col: int | None = None
    desc_col: int | None = None
    sheet: str | None = None


@dataclass(frozen=True)
class SimulationResult:
    path: Path
    manifest_path: Path
    manifest: dict[str, Any]


class _Row:
    """One row of the output, remembering where it came from."""

    def __init__(self, source_index: int | None, cells: list[Any]):
        self.source_index = source_index  # 0-based index in the source sheet, None if added
        self.cells = cells


# -- reading the source ---------------------------------------------------------------


def _best(roles: list[tuple[int, ColumnRole, float]], wanted: ColumnRole) -> int | None:
    candidates = [(score, -index) for index, role, score in roles if role is wanted]
    return None if not candidates else -max(candidates)[1]


def _price_of(cell: Any, column: ColumnProfile) -> Decimal | None:
    if isinstance(cell, bool) or cell is None:
        return None
    if isinstance(cell, int | float | Decimal):
        number = Decimal(repr(cell)) if isinstance(cell, float) else Decimal(cell)
        return number if number.is_finite() else None
    if not isinstance(cell, str):
        return None
    fmt = column.number_format
    if fmt is not None and fmt.ambiguous:
        return None
    try:
        return parse_number(cell, (fmt.decimal_sep if fmt else None) or ".", (fmt.thousands_sep if fmt else None) or "")
    except CellError:
        return None


def _format_number(value: Decimal, decimals: int, decimal_sep: str, thousands_sep: str, grouped: bool) -> str:
    text = format(value.quantize(Decimal(1).scaleb(-decimals), rounding=ROUND_HALF_UP), "f")
    whole, _, fraction = text.partition(".")
    if grouped and thousands_sep:
        whole = f"{int(whole):,}".replace(",", thousands_sep)
    return whole + (decimal_sep + fraction if fraction else "")


def _with_price(cell: Any, new: Decimal, decimals: int, column: ColumnProfile) -> Any:
    """`new` written like the cell it replaces: a number stays a number, text keeps its format."""
    if not isinstance(cell, str):
        return new
    fmt = column.number_format
    decimal_sep = (fmt.decimal_sep if fmt else None) or "."
    thousands_sep = (fmt.thousands_sep if fmt else None) or ""

    def swap(match: re.Match) -> str:
        grouped = bool(thousands_sep) and thousands_sep in match.group(0)
        return _format_number(new, decimals, decimal_sep, thousands_sep, grouped)

    return _NUMBER_TOKEN.sub(swap, cell, count=1)


def _decimals(value: Decimal) -> int:
    return max(2, -value.as_tuple().exponent) if isinstance(value.as_tuple().exponent, int) else 2


def _text(cell: Any) -> str:
    return "" if cell is None else str(cell).strip()


# -- the simulation ---------------------------------------------------------------------


def _pct(old: Decimal, new: Decimal) -> str:
    return str(((new - old) / old * 100).quantize(_CENT, rounding=ROUND_HALF_UP))


def _new_code(sample: Any, taken: set[str], numeric_max: int | None, serial: int) -> Any:
    """A code that is not in the list, shaped like the one it was cloned from."""
    text = _text(sample)
    if isinstance(sample, int | Decimal) and not isinstance(sample, bool) or text.isdigit():
        top = max(numeric_max or 0, 0) + serial
        while str(top) in taken:
            top += 1
        if isinstance(sample, int | Decimal) and not isinstance(sample, bool):
            return int(top)
        return str(top).zfill(len(text))
    code = f"{text}-SIM{serial}"
    while code in taken:
        code += "x"
    return code


def _plan(source_rows, profile, columns, options: SimulationOptions):
    """Decide, with the seed, which rows change and how. Returns the plan and the eligible rows."""
    code_i, price_i = columns["code"], columns["price"]
    price_col = profile.columns[price_i]
    data = profile.indexes_of(RowClass.DATA)
    codes: dict[str, int] = {}
    for index in data:
        code = _text(source_rows[index][code_i]) if code_i < len(source_rows[index]) else ""
        if code:
            codes[code] = codes.get(code, 0) + 1
    eligible = []
    for index in data:
        row = source_rows[index]
        code = _text(row[code_i]) if code_i < len(row) else ""
        price = _price_of(row[price_i] if price_i < len(row) else None, price_col)
        if code and codes[code] == 1 and price is not None and price >= options.min_price:
            eligible.append(index)
    rng = random.Random(options.seed)
    pool = list(eligible)
    anomaly = None
    if options.anomaly and pool:
        anomaly = pool.pop(rng.randrange(len(pool)))
    removed = sorted(rng.sample(pool, min(options.removed, len(pool))))
    pool = [index for index in pool if index not in set(removed)]
    drift_count = int((Decimal(len(pool)) * options.drift_share).to_integral_value(rounding=ROUND_HALF_UP))
    drifting = sorted(rng.sample(pool, min(drift_count, len(pool))))
    drifts = {}
    span = options.drift_max - options.drift_min
    for index in drifting:
        down = rng.random() < float(options.down_share)
        ratio = options.drift_min + Decimal(str(round(rng.random(), 6))) * span
        drifts[index] = -ratio if down else ratio
    clones = [eligible[rng.randrange(len(eligible))] for _ in range(options.added)] if eligible else []
    factors = [Decimal(str(round(0.8 + rng.random() * 0.4, 4))) for _ in clones]
    return eligible, anomaly, removed, drifts, list(zip(clones, factors, strict=True))


def _detect(profile, options: SimulationOptions) -> dict[str, int | None]:
    roles = [
        (item.index, item.role, item.score)
        for item in HeuristicColumnClassifier().classify_sheet(profile)
    ]
    columns = {
        "code": options.code_col if options.code_col is not None else _best(roles, ColumnRole.CODE),
        "price": options.price_col if options.price_col is not None else _best(roles, ColumnRole.COST_PRICE),
        "desc": options.desc_col if options.desc_col is not None else _best(roles, ColumnRole.DESCRIPTION),
    }
    if columns["code"] is None:
        raise ValueError("no code column was detected; pass --code-col")
    if columns["price"] is None:
        raise ValueError("no price column was detected; pass --price-col")
    return columns


def _sheet_title(name: str) -> str:
    title = _BAD_SHEET_CHARS.sub("_", name).strip()[:SHEET_NAME_LIMIT] or "LIST"
    return "LIST" if title.upper() == MARKER else title


def simulate(source: Path, out_dir: Path, options: SimulationOptions) -> SimulationResult:
    source, out_dir = Path(source), Path(out_dir)
    if out_dir.resolve() == source.resolve().parent:
        raise ValueError("refusing to write simulated files next to the real source; use another folder")
    resolved = out_dir.resolve()
    if resolved.is_relative_to(REPO_ROOT) and not resolved.is_relative_to(RAW_ROOT):
        raise ValueError("the output folder must be outside the repository or under data/raw")

    sheet = FileSpreadsheetReader().read_sheet(source, options.sheet)
    source_rows = [tuple(row) for row in sheet.rows]
    profile = profile_sheet(source_rows)
    if profile.header_row is None:
        raise ValueError("the source sheet has no detectable header row")
    columns = _detect(profile, options)
    code_i, price_i = columns["code"], columns["price"]
    price_col = profile.columns[price_i]
    width = max(len(row) for row in source_rows)
    header_row = source_rows[profile.header_row]

    eligible, anomaly, removed, drifts, clones = _plan(source_rows, profile, columns, options)
    rows = [_Row(index, list(row) + [None] * (width - len(row))) for index, row in enumerate(source_rows)]
    by_source = {row.source_index: row for row in rows}
    old_prices: dict[int, Decimal] = {}
    new_prices: dict[int, Decimal] = {}

    def reprice(index: int, factor: Decimal) -> None:
        row = by_source[index]
        old = _price_of(row.cells[price_i], price_col)
        decimals = _decimals(old)
        new = (old * factor).quantize(Decimal(1).scaleb(-decimals), rounding=ROUND_HALF_UP)
        if new == old or new <= 0:
            return
        old_prices[index], new_prices[index] = old, new
        row.cells[price_i] = _with_price(row.cells[price_i], new, decimals, price_col)

    for index, ratio in drifts.items():
        reprice(index, 1 + ratio)
    if anomaly is not None:
        reprice(anomaly, options.anomaly_factor)

    # added rows go after the last data row that is still there
    taken = {_text(row[code_i]) for row in source_rows if code_i < len(row)}
    numeric = [int(code) for code in taken if code.isdigit()]
    numeric_max = max(numeric, default=None)
    removed_set = set(removed)
    last_data = max(profile.indexes_of(RowClass.DATA))
    new_rows: list[_Row] = []
    for serial, (clone, factor) in enumerate(clones, start=1):
        base = by_source[clone]
        cells = list(base.cells)
        code = _new_code(cells[code_i], taken, numeric_max, serial)
        taken.add(_text(code))
        cells[code_i] = code
        if columns["desc"] is not None:
            cells[columns["desc"]] = f"{MARKER} NEW ITEM {serial}"
        old = _price_of(base.cells[price_i], price_col)
        decimals = _decimals(old)
        new = (old * factor).quantize(Decimal(1).scaleb(-decimals), rounding=ROUND_HALF_UP)
        cells[price_i] = _with_price(base.cells[price_i], new, decimals, price_col)
        new_rows.append(_Row(None, cells))
    kept = [row for row in rows if row.source_index not in removed_set]
    anchor = max(position for position, row in enumerate(kept) if (row.source_index or 0) <= last_data)
    kept[anchor + 1 : anchor + 1] = new_rows

    moved = None
    if options.move_column:
        target = price_i + 1 if price_i + 1 < width else price_i - 1
        for row in kept:
            row.cells[price_i], row.cells[target] = row.cells[target], row.cells[price_i]
        moved = {"header": _text(header_row[price_i]), "from_index": price_i, "to_index": target}

    def position(row: _Row) -> int:
        return kept.index(row) + 1

    def change_entry(index: int) -> dict[str, Any]:
        row = by_source[index]
        return {
            "code": _text(source_rows[index][code_i]),
            "row": position(row),
            "source_row": index + 1,
            "old": str(old_prices[index]),
            "new": str(new_prices[index]),
            "pct": _pct(old_prices[index], new_prices[index]),
        }

    changes = [change_entry(index) for index in sorted(new_prices) if index != anomaly]
    anomaly_entry = change_entry(anomaly) if anomaly in new_prices else None
    desc_i = columns["desc"]
    manifest = {
        "simulated": True,
        "tool": TOOL,
        "note": "Generated demo data. It is not a real supplier price list.",
        "seed": options.seed,
        "source": {
            "file": source.name,
            "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "sheet": sheet.name,
            "rows": len(source_rows),
        },
        "options": {key: str(value) if isinstance(value, Decimal) else value for key, value in asdict(options).items()},
        "columns": {
            "code_index": code_i,
            "code_header": _text(header_row[code_i]),
            "price_index": price_i,
            "price_header": _text(header_row[price_i]),
        },
        "price_changes": changes,
        "added": [
            {
                "code": _text(row.cells[code_i]),
                "row": position(row),
                "description": "" if desc_i is None else _text(row.cells[desc_i]),
                "price": str(_price_of(row.cells[price_i if moved is None else moved["to_index"]], price_col) or ""),
            }
            for row in new_rows
        ],
        "removed": [
            {
                "code": _text(source_rows[index][code_i]),
                "source_row": index + 1,
                "description": "" if desc_i is None else _text(source_rows[index][desc_i]),
            }
            for index in removed
        ],
        "moved_column": moved,
        "anomaly": anomaly_entry,
    }
    manifest["summary"] = {
        "price_changes": len(changes),
        "increases": sum(1 for item in changes if Decimal(item["pct"]) > 0),
        "decreases": sum(1 for item in changes if Decimal(item["pct"]) < 0),
        "added": len(new_rows),
        "removed": len(removed),
        "column_moved": moved is not None,
        "anomaly": anomaly_entry is not None,
        # priced products present in both files with the same price
        "unchanged": len(eligible) - len(removed) - len(changes) - (anomaly_entry is not None),
    }

    out_dir.mkdir(parents=True, exist_ok=True)
    stem = f"{source.stem}-{MARKER}-week2"
    path, manifest_path = out_dir / f"{stem}.xlsx", out_dir / f"{stem}.manifest.json"
    manifest["output"] = {"file": path.name, "sheet": _sheet_title(sheet.name), "rows": len(kept)}
    _write_workbook(path, sheet.name, [row.cells for row in kept], source.name, options.seed)
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return SimulationResult(path, manifest_path, manifest)


def _write_workbook(path: Path, sheet_name: str, rows: list[list[Any]], source_name: str, seed: int) -> None:
    book = Workbook()
    main_sheet = book.active
    main_sheet.title = _sheet_title(sheet_name)
    for r, cells in enumerate(rows, start=1):
        for c, value in enumerate(cells, start=1):
            if value is None:
                continue
            cell = main_sheet.cell(row=r, column=c, value=value)
            if isinstance(value, str):
                cell.data_type = "s"  # untrusted text must never become a formula
    note = book.create_sheet(MARKER)
    note["A1"] = f"{MARKER} DATA: generated by {TOOL} from {source_name} (seed {seed})."
    note["A2"] = "This is not a real supplier price list. See the manifest next to this file."
    for cell in (note["A1"], note["A2"]):
        cell.data_type = "s"
    book.properties.title = f"{MARKER} week 2 of {source_name}"
    book.properties.subject = f"{MARKER} DATA"
    book.properties.description = "Generated demo data, not a real supplier price list."
    book.properties.creator = TOOL
    book.properties.created = book.properties.modified = _FIXED_STAMP
    book.save(path)


# -- command line ---------------------------------------------------------------------------


def _default_out_dir(source: Path) -> Path | None:
    parent = source.resolve().parent
    if parent.parent.parent.name == "raw":
        return parent.with_name(parent.name + "-simulated")
    return None


def _ratio(text: str) -> Decimal:
    value = Decimal(text)
    if not value.is_finite() or value < 0:
        raise argparse.ArgumentTypeError(f"{text!r} is not a ratio of zero or more")
    return value


def summary_text(result: SimulationResult) -> str:
    summary = result.manifest["summary"]
    lines = [
        f"{MARKER} week 2 written: {result.path}",
        f"manifest: {result.manifest_path}",
        f"price changes: {summary['price_changes']} "
        f"({summary['increases']} up, {summary['decreases']} down)",
        f"added: {summary['added']}",
        f"removed: {summary['removed']}",
        f"column moved: {result.manifest['moved_column']}",
        f"planted anomaly: {result.manifest['anomaly']}",
    ]
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    defaults = SimulationOptions()
    parser.add_argument("source", type=Path, help="a real snapshot (.xlsx, .xls or .csv)")
    parser.add_argument("--out-dir", type=Path)
    parser.add_argument("--seed", type=int, default=defaults.seed)
    parser.add_argument("--drift-share", type=_ratio, default=defaults.drift_share)
    parser.add_argument("--drift-min", type=_ratio, default=defaults.drift_min)
    parser.add_argument("--drift-max", type=_ratio, default=defaults.drift_max)
    parser.add_argument("--down-share", type=_ratio, default=defaults.down_share)
    parser.add_argument("--added", type=int, default=defaults.added)
    parser.add_argument("--removed", type=int, default=defaults.removed)
    parser.add_argument("--move-column", action="store_true")
    parser.add_argument("--anomaly", action="store_true")
    parser.add_argument("--anomaly-factor", type=_ratio, default=defaults.anomaly_factor)
    parser.add_argument("--min-price", type=_ratio, default=defaults.min_price)
    parser.add_argument("--code-col", type=int)
    parser.add_argument("--price-col", type=int)
    parser.add_argument("--desc-col", type=int)
    parser.add_argument("--sheet")
    args = parser.parse_args(argv)

    out_dir = args.out_dir or _default_out_dir(args.source)
    if out_dir is None:
        parser.error("--out-dir is required unless the source is under data/raw/<id>/<date>/")
    if args.drift_min > args.drift_max:
        parser.error("--drift-min cannot be larger than --drift-max")
    options = replace(
        defaults,
        seed=args.seed,
        drift_share=args.drift_share,
        drift_min=args.drift_min,
        drift_max=args.drift_max,
        down_share=args.down_share,
        added=args.added,
        removed=args.removed,
        move_column=args.move_column,
        anomaly=args.anomaly,
        anomaly_factor=args.anomaly_factor,
        min_price=args.min_price,
        code_col=args.code_col,
        price_col=args.price_col,
        desc_col=args.desc_col,
        sheet=args.sheet,
    )
    try:
        result = simulate(args.source, out_dir, options)
    except (ValueError, OSError) as error:
        print(f"error: {error}", file=sys.stderr)
        return 1
    print(summary_text(result))
    return 0


if __name__ == "__main__":
    sys.exit(main())
