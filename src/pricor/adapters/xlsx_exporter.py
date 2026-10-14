"""XLSX exporter (openpyxl).

The workbook has a `Products` sheet (contract columns + traceability) and a
`Rejected rows` sheet. Text is always stored as a text cell, never as a
formula, so a supplier's `=...` cannot run; codes, barcodes and original cells
keep leading zeros. Money and rates are numeric cells whose number format shows
their own decimals. A spreadsheet number is an IEEE double, exact up to 15
significant digits: a decimal with more digits is written as text instead of
being rounded silently.
"""

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from pricor.adapters.export_common import (
    has_conversion,
    product_columns,
    product_values,
    rejected_header,
    rejected_width,
)
from pricor.ports.exporter import ImportExport

PRODUCTS_SHEET = "Products"
REJECTED_SHEET = "Rejected rows"
MAX_EXACT_DIGITS = 15
_TEXT_COLUMNS = frozenset(
    {"supplier_code", "description", "currency", "barcode", "fx_source", "fx_date"}
)


def _set_text(sheet: Worksheet, row: int, column: int, text: str) -> None:
    cell = sheet.cell(row=row, column=column, value=text)
    cell.data_type = "s"  # assigning "=..." would otherwise make it a formula
    cell.number_format = "@"


def _set_number(sheet: Worksheet, row: int, column: int, value: Decimal | int) -> None:
    if isinstance(value, Decimal) and len(value.normalize().as_tuple().digits) > MAX_EXACT_DIGITS:
        _set_text(sheet, row, column, format(value, "f"))
        return
    cell = sheet.cell(row=row, column=column, value=value)
    decimals = -value.as_tuple().exponent if isinstance(value, Decimal) else 0
    cell.number_format = "0." + "0" * decimals if decimals > 0 else "0"


def _header(sheet: Worksheet, names: list[str]) -> None:
    for column, name in enumerate(names, start=1):
        _set_text(sheet, 1, column, name)
        sheet.cell(row=1, column=column).font = Font(bold=True)
    sheet.freeze_panes = "A2"


def _products(sheet: Worksheet, data: ImportExport) -> None:
    names = product_columns(data)
    converted = has_conversion(data)
    _header(sheet, list(names))
    for row, item in enumerate(data.accepted, start=2):
        values = product_values(item.row_number, item.record, converted)
        for column, (name, value) in enumerate(zip(names, values, strict=True), start=1):
            if value is None:
                continue
            if name in _TEXT_COLUMNS:
                _set_text(sheet, row, column, value)
            else:
                _set_number(sheet, row, column, value)


def _rejected(sheet: Worksheet, data: ImportExport) -> None:
    _header(sheet, rejected_header(data))
    width = rejected_width(data)
    for row, entry in enumerate(data.rejected, start=2):
        sheet.cell(row=row, column=1, value=entry.row_number)
        _set_text(sheet, row, 2, "; ".join(entry.reasons))
        for offset in range(width):
            if offset < len(entry.cells) and entry.cells[offset] != "":
                _set_text(sheet, row, 3 + offset, entry.cells[offset])


def _save(book: Workbook) -> bytes:
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


class XlsxExporter:
    format = "xlsx"
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    extension = ".xlsx"

    def export(self, data: ImportExport) -> bytes:
        book = Workbook()
        products = book.active
        products.title = PRODUCTS_SHEET
        _products(products, data)
        _rejected(book.create_sheet(REJECTED_SHEET), data)
        return _save(book)

    def export_rejected(self, data: ImportExport) -> bytes:
        book = Workbook()
        report = book.active
        report.title = REJECTED_SHEET
        _rejected(report, data)
        return _save(book)
