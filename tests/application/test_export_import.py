import csv
import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from pricor.adapters.csv_exporter import CsvExporter
from pricor.adapters.xlsx_exporter import XlsxExporter
from pricor.application.errors import InvalidInputError, InvalidOperationError
from pricor.application.export_import import ExportImport
from pricor.domain.cells import display
from pricor.ports.import_session_repository import SessionNotFoundError

from ..persistence.sessions import Flow


def setup(supplier="Ferretería Ejemplo, S.A./Sur"):
    flow = Flow(supplier=supplier)
    export = ExportImport(flow.use_case.get, [CsvExporter(), XlsxExporter()])
    return flow, export


def rows_of(file):
    return list(csv.reader(io.StringIO(file.content.decode("utf-8-sig"), newline="")))


def test_products_export_matches_the_approved_result():
    flow, export = setup()
    session = flow.approved()
    file = export.products(session.id, "csv")
    rows = rows_of(file)
    accepted = session.approved_result.accepted
    assert len(rows) - 1 == len(accepted)
    assert [row[0] for row in rows[1:]] == [item.record.supplier_code for item in accepted]
    assert [int(row[8]) for row in rows[1:]] == [item.row_number for item in accepted]
    assert "00123" in [row[0] for row in rows[1:]]  # leading zero kept
    assert [Decimal(row[2]) for row in rows[1:]] == [item.record.unit_cost_net for item in accepted]
    assert file.media_type.startswith("text/csv")


def test_rejected_report_has_one_line_per_rejected_row_with_original_cells():
    flow, export = setup()
    session = flow.approved()
    rejected = session.approved_result.rejected
    assert rejected  # the sample file has a bad price and a datetime row
    rows = rows_of(export.rejected(session.id, "csv"))
    assert rows[0][:2] == ["source_row", "reasons"]
    assert rows[0][2:] == ["Codigo", "Descripcion", "Precio lista", "IVA"]  # the file's own headers
    assert len(rows) - 1 == len(rejected)
    by_row = {int(row[0]): row for row in rows[1:]}
    for item in rejected:
        line = by_row[item.row_number]
        assert line[1] == "; ".join(item.reasons)
        assert line[2:6] == [("" if c is None else display(c)) for c in item.cells][:4]
    assert by_row[8][2:] == ["A6", "Brida", "abc", "21%"]


def test_xlsx_products_and_rejected_sheet_agree_with_the_session():
    flow, export = setup()
    session = flow.approved()
    book = load_workbook(io.BytesIO(export.products(session.id, "xlsx").content))
    assert book["Products"].max_row - 1 == len(session.approved_result.accepted)
    assert book["Rejected rows"].max_row - 1 == len(session.approved_result.rejected)
    only = load_workbook(io.BytesIO(export.rejected(session.id, "xlsx").content))
    assert only.sheetnames == ["Rejected rows"]


def test_filenames_are_safe_and_identify_the_import():
    flow, export = setup()
    session = flow.approved()
    products = export.products(session.id, "csv")
    rejected = export.rejected(session.id, "xlsx")
    assert products.filename == f"pricor-ferreteria-ejemplo-s-a-sur-{session.id[:8]}-products.csv"
    assert rejected.filename == f"pricor-ferreteria-ejemplo-s-a-sur-{session.id[:8]}-rejected.xlsx"
    hostile = Flow(supplier='../"\r\nX-Evil: 1')
    odd = hostile.approved()
    name = ExportImport(hostile.use_case.get, [CsvExporter()]).products(odd.id, "csv").filename
    assert name.isascii() and all(ch.isalnum() or ch in "-." for ch in name)


def test_sessions_that_are_not_approved_cannot_be_exported():
    flow, export = setup()
    session = flow.ready()
    with pytest.raises(InvalidOperationError):
        export.products(session.id, "csv")
    with pytest.raises(InvalidOperationError):
        export.rejected(session.id, "xlsx")


def test_unknown_formats_and_sessions_are_refused():
    flow, export = setup()
    session = flow.approved()
    with pytest.raises(InvalidInputError):
        export.products(session.id, "pdf")
    with pytest.raises(SessionNotFoundError):
        export.products("missing", "csv")
