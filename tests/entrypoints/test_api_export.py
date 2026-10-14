import csv
import io
from decimal import Decimal

from openpyxl import load_workbook

from .test_api import ROWS, answer, make_client, upload

ROWS_WITH_ZEROS = [
    ["Codigo", "Descripcion", "Precio lista", "IVA"],
    ["00123", "Tornillo", "1.210,00", "21%"],
    ["00456", "Tuerca", "968,00", "21%"],
    ["A9", "Brida", "abc", "21%"],
    ["", "Sin codigo", "10,00", "21%"],
]


def approved_import(client, rows=ROWS):
    state = upload(client, rows=rows).json()
    session_id = state["id"]
    answers = {"number_format": "decimal_comma", "price_includes_vat": "yes", "currency": "ARS"}
    for question in state["questions"]:
        answer(client, session_id, question["id"], answers[question["id"]])
    assert client.post(f"/imports/{session_id}/approve").status_code == 200
    return session_id


def csv_rows(response):
    return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"), newline="")))


def test_export_requires_an_approved_session():
    client = make_client()
    session_id = upload(client).json()["id"]
    answer(client, session_id, "price_includes_vat", "yes")
    answer(client, session_id, "currency", "ARS")  # ready for preview, not approved
    for path in ("export", "rejected"):
        for fmt in ("csv", "xlsx"):
            response = client.get(f"/imports/{session_id}/{path}", params={"format": fmt})
            assert response.status_code == 409
            assert response.json()["error"]["code"] == "invalid_operation"


def test_unknown_session_is_a_404_and_unknown_format_a_422():
    client = make_client()
    assert client.get("/imports/nope/export", params={"format": "csv"}).status_code == 404
    session_id = approved_import(client)
    response = client.get(f"/imports/{session_id}/export", params={"format": "pdf"})
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "invalid_input"


def test_csv_export_has_the_right_headers_and_every_accepted_row():
    client = make_client()
    session_id = approved_import(client)
    response = client.get(f"/imports/{session_id}/export", params={"format": "csv"})
    assert response.status_code == 200
    assert response.headers["content-type"] == "text/csv; charset=utf-8"
    disposition = response.headers["content-disposition"]
    assert disposition == f'attachment; filename="pricor-ferreteria-ejemplo-{session_id[:8]}-products.csv"'
    rows = csv_rows(response)
    totals = client.get(f"/imports/{session_id}/preview").json()["totals"]
    assert len(rows) - 1 == totals["accepted"] == 5
    assert rows[0][:4] == ["supplier_code", "description", "unit_cost_net", "currency"]


def test_format_defaults_to_csv():
    client = make_client()
    session_id = approved_import(client)
    response = client.get(f"/imports/{session_id}/export")
    assert response.headers["content-type"].startswith("text/csv")


def test_rejected_report_count_equals_the_sessions_rejected_count():
    client = make_client()
    session_id = approved_import(client, ROWS_WITH_ZEROS)
    totals = client.get(f"/imports/{session_id}/preview").json()["totals"]
    for fmt in ("csv", "xlsx"):
        response = client.get(f"/imports/{session_id}/rejected", params={"format": fmt})
        assert response.status_code == 200
        assert f"-rejected.{fmt}" in response.headers["content-disposition"]
        if fmt == "csv":
            rows = csv_rows(response)
            assert [row[0] for row in rows[1:]] == [str(r["row_number"]) for r in totals["rejected_rows"]]
        else:
            sheet = load_workbook(io.BytesIO(response.content))["Rejected rows"]
            assert sheet.max_row - 1 == totals["rejected"] == 2


def test_xlsx_export_keeps_leading_zeros_and_exact_decimals():
    client = make_client()
    session_id = approved_import(client, ROWS_WITH_ZEROS)
    response = client.get(f"/imports/{session_id}/export", params={"format": "xlsx"})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    book = load_workbook(io.BytesIO(response.content))
    products = book["Products"]
    assert [products["A2"].value, products["A3"].value] == ["00123", "00456"]
    assert Decimal(str(products["C2"].value)) == Decimal("1000.00")  # 1210,00 without 21% VAT
    assert Decimal(str(products["C3"].value)) == Decimal("800.00")
    assert book["Rejected rows"].max_row == 3


def test_csv_export_keeps_leading_zeros_as_text():
    client = make_client()
    session_id = approved_import(client, ROWS_WITH_ZEROS)
    rows = csv_rows(client.get(f"/imports/{session_id}/export", params={"format": "csv"}))
    assert [row[0] for row in rows[1:]] == ["00123", "00456"]
    assert rows[1][2] == "1000.00"
