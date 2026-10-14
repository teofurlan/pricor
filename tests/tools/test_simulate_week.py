import hashlib
import json
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from tools.corpus.simulate_week import SimulationOptions, main, simulate

HEADER = ["Codigo", "Descripcion", "Rubro", "Precio", "Moneda"]
CODE, PRICE = 0, 3


def build_source(path: Path, rows=None, text_prices=False) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "LISTA"
    sheet.append(["LISTA DE PRECIOS OCTUBRE", None, None, None, None])
    sheet.append(HEADER)
    for number in range(1, 41):
        price = f"$ {1000 + number * 37},50" if text_prices else Decimal(100) + Decimal("7.5") * number
        sheet.append([f"P{number:03d}", f"Producto {number}", "Rubro A", price, "$"])
    sheet.append(["", "", "", None, ""])  # empty row
    sheet.append(["TOTAL", None, None, 99999, None])
    for extra in rows or []:
        sheet.append(extra)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path


@pytest.fixture
def source(tmp_path):
    return build_source(tmp_path / "in" / "lista.xlsx")


def options(**overrides) -> SimulationOptions:
    values = {"seed": 7, "code_col": CODE, "price_col": PRICE, "desc_col": 1}
    values.update(overrides)
    return SimulationOptions(**values)


def table(path: Path, sheet=None) -> list[list]:
    book = load_workbook(path, data_only=True)
    return [list(row) for row in (book[sheet] if sheet else book.worksheets[0]).iter_rows(values_only=True)]


def products(rows: list[list], code_index=CODE, price_index=PRICE) -> dict[str, Decimal]:
    return {
        str(row[code_index]): Decimal(str(row[price_index]))
        for row in rows
        if row[code_index] and str(row[code_index]).startswith("P") and row[price_index] is not None
    }


def run(source, tmp_path, **overrides):
    return simulate(source, tmp_path / "out", options(**overrides))


def test_the_output_is_labeled_simulated_in_the_name_a_sheet_and_the_file_properties(source, tmp_path):
    result = run(source, tmp_path)

    assert "SIMULATED" in result.path.name and result.path.suffix == ".xlsx"
    book = load_workbook(result.path)
    assert "SIMULATED" in book.sheetnames and book.worksheets[0].title == "LISTA"
    assert "SIMULATED" in book["SIMULATED"]["A1"].value
    assert "SIMULATED" in (book.properties.title or "")
    assert result.manifest_path.parent == result.path.parent
    assert result.manifest["simulated"] is True


def test_the_real_file_is_never_modified_and_never_overwritten(source, tmp_path):
    before = hashlib.sha256(source.read_bytes()).hexdigest()

    result = run(source, tmp_path)

    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    assert result.manifest["source"]["sha256"] == before
    with pytest.raises(ValueError, match="source"):
        simulate(source, source.parent, options())


def test_the_manifest_matches_the_actual_difference_between_the_two_files(source, tmp_path):
    result = run(source, tmp_path, added=4, removed=3, anomaly=True)
    manifest = result.manifest
    old, new = products(table(source)), products(table(result.path))

    removed = {item["code"] for item in manifest["removed"]}
    added = {item["code"] for item in manifest["added"]}
    changed = {item["code"]: item for item in manifest["price_changes"]}
    anomaly = manifest["anomaly"]
    assert set(old) - set(new) == removed
    assert set(new) - set(old) == added
    actual = {code for code in set(old) & set(new) if old[code] != new[code]}
    assert actual == set(changed) | {anomaly["code"]}
    for code, item in changed.items():
        assert (Decimal(item["old"]), Decimal(item["new"])) == (old[code], new[code])
        assert Decimal(item["pct"]) == ((new[code] - old[code]) / old[code] * 100).quantize(Decimal("0.01"))
    assert (Decimal(anomaly["old"]), Decimal(anomaly["new"])) == (old[anomaly["code"]], new[anomaly["code"]])
    summary = manifest["summary"]
    assert summary["added"] == 4 and summary["removed"] == 3 and summary["price_changes"] == len(changed)
    assert summary["increases"] + summary["decreases"] == len(changed)
    assert summary["unchanged"] == len(set(old) & set(new)) - len(changed) - 1


def test_manifest_rows_point_at_the_real_rows_of_each_file(source, tmp_path):
    result = run(source, tmp_path, added=2, removed=2)
    old_rows, new_rows = table(source), table(result.path)

    for item in result.manifest["price_changes"]:
        assert old_rows[item["source_row"] - 1][CODE] == item["code"]
        assert new_rows[item["row"] - 1][CODE] == item["code"]
    for item in result.manifest["removed"]:
        assert old_rows[item["source_row"] - 1][CODE] == item["code"]
    for item in result.manifest["added"]:
        assert new_rows[item["row"] - 1][CODE] == item["code"]


def test_the_same_seed_gives_the_same_simulation_and_another_seed_a_different_one(source, tmp_path):
    first = simulate(source, tmp_path / "a", options(seed=3, added=3, removed=2, anomaly=True))
    again = simulate(source, tmp_path / "b", options(seed=3, added=3, removed=2, anomaly=True))
    other = simulate(source, tmp_path / "c", options(seed=4, added=3, removed=2, anomaly=True))

    assert first.manifest == again.manifest
    assert table(first.path) == table(again.path)
    assert first.manifest_path.read_text(encoding="utf-8") == again.manifest_path.read_text(encoding="utf-8")
    assert first.manifest["price_changes"] != other.manifest["price_changes"]


def test_price_changes_stay_inside_the_requested_drift_range(source, tmp_path):
    result = run(source, tmp_path, drift_min=Decimal("0.02"), drift_max=Decimal("0.05"), drift_share=Decimal("1"))

    pcts = [abs(Decimal(item["pct"])) for item in result.manifest["price_changes"]]
    assert pcts and all(Decimal("1.9") <= pct <= Decimal("5.1") for pct in pcts)  # rounding to cents moves it a hair
    assert {Decimal(item["pct"]) > 0 for item in result.manifest["price_changes"]} == {True, False}


def test_the_planted_anomaly_is_beyond_the_default_threshold_and_untouched_by_drift(source, tmp_path):
    result = run(source, tmp_path, anomaly=True, drift_share=Decimal("1"))

    anomaly = result.manifest["anomaly"]
    assert Decimal(anomaly["pct"]) == Decimal("200.00")
    assert anomaly["code"] not in {item["code"] for item in result.manifest["price_changes"]}
    assert anomaly["code"] not in {item["code"] for item in result.manifest["removed"]}


def test_no_anomaly_unless_asked(source, tmp_path):
    assert run(source, tmp_path).manifest["anomaly"] is None


def test_added_products_are_new_unique_codes_and_removed_ones_are_gone(source, tmp_path):
    result = run(source, tmp_path, added=5, removed=4)
    new = table(result.path)

    codes = [row[CODE] for row in new if row[CODE]]
    assert len(codes) == len(set(codes))
    added = [item["code"] for item in result.manifest["added"]]
    assert len(added) == 5 and not set(added) & set(products(table(source)))
    assert all("SIMULATED" in next(row for row in new if row[CODE] == code)[1] for code in added)
    assert len(result.manifest["removed"]) == 4


def test_moving_a_column_shifts_every_row_and_is_recorded(source, tmp_path):
    result = run(source, tmp_path, move_column=True)

    new = table(result.path)
    assert new[1] == ["Codigo", "Descripcion", "Rubro", "Moneda", "Precio"]
    assert result.manifest["moved_column"] == {"header": "Precio", "from_index": 3, "to_index": 4}
    assert new[2][3:] == ["$", Decimal("107.5")] or new[2][3:] == ["$", 107.5]


def test_without_options_the_list_only_gets_the_label(source, tmp_path):
    result = run(source, tmp_path, drift_share=Decimal("0"), added=0, removed=0)

    assert table(result.path) == table(source)
    assert result.manifest["price_changes"] == [] and result.manifest["added"] == []
    assert result.manifest["removed"] == [] and result.manifest["moved_column"] is None


def test_preamble_total_and_unpriced_rows_are_never_touched(tmp_path):
    source = build_source(tmp_path / "in" / "lista.xlsx", rows=[["P999", "Sin precio", "X", 0, "$"]])

    result = run(source, tmp_path, drift_share=Decimal("1"), removed=5, added=2)

    new = table(result.path)
    assert new[0][0] == "LISTA DE PRECIOS OCTUBRE"
    assert any(row[CODE] == "TOTAL" and row[PRICE] == 99999 for row in new)
    assert any(row[CODE] == "P999" and row[PRICE] == 0 for row in new)


def test_text_prices_keep_their_own_format(tmp_path):
    source = build_source(tmp_path / "in" / "texto.xlsx", text_prices=True)

    result = run(source, tmp_path, drift_share=Decimal("1"))

    for item in result.manifest["price_changes"]:
        row = next(r for r in table(result.path) if r[CODE] == item["code"])
        assert str(row[PRICE]).startswith("$ ") and str(row[PRICE]).count(",") == 1
        assert Decimal(str(row[PRICE]).removeprefix("$ ").replace(".", "").replace(",", ".")) == Decimal(item["new"])


def test_columns_are_detected_when_not_given(source, tmp_path):
    result = simulate(source, tmp_path / "out", SimulationOptions(seed=1, added=1))

    assert (result.manifest["columns"]["code_header"], result.manifest["columns"]["price_header"]) == (
        "Codigo",
        "Precio",
    )


def test_an_undetectable_price_column_is_an_error_not_a_guess(tmp_path):
    path = tmp_path / "in" / "raro.xlsx"
    path.parent.mkdir()
    book = Workbook()
    book.active.append(["a", "b"])
    book.active.append(["x", "y"])
    book.save(path)

    with pytest.raises(ValueError, match="price"):
        simulate(path, tmp_path / "out", SimulationOptions(seed=1, code_col=0))


def test_the_command_writes_the_files_and_prints_a_summary(source, tmp_path, capsys):
    code = main([str(source), "--out-dir", str(tmp_path / "cli"), "--seed", "5", "--added", "2", "--removed", "1",
                 "--anomaly", "--code-col", "0", "--price-col", "3"])

    assert code == 0
    out = capsys.readouterr().out
    assert "SIMULATED" in out and "added: 2" in out and "removed: 1" in out
    written = sorted(path.name for path in (tmp_path / "cli").iterdir())
    assert len(written) == 2 and any(name.endswith(".manifest.json") for name in written)
    manifest = json.loads(next((tmp_path / "cli").glob("*.manifest.json")).read_text(encoding="utf-8"))
    assert manifest["options"]["seed"] == 5


def test_the_output_folder_defaults_next_to_a_raw_snapshot(tmp_path, capsys):
    raw = build_source(tmp_path / "data" / "raw" / "emilio" / "2026-10-13" / "lista.xlsx")

    assert main([str(raw), "--code-col", "0", "--price-col", "3"]) == 0

    assert (tmp_path / "data" / "raw" / "emilio" / "2026-10-13-simulated").is_dir()


def test_without_a_raw_snapshot_path_an_output_folder_is_required(source, capsys):
    with pytest.raises(SystemExit):
        main([str(source)])


def test_files_are_never_written_inside_the_repository_outside_the_ignored_raw_folder(source):
    inside = Path(__file__).resolve().parents[2] / "docs" / "simulated-output"

    with pytest.raises(ValueError, match="outside the repository"):
        simulate(source, inside, options())
    assert not inside.exists()


def test_products_cheaper_than_the_minimum_price_are_left_alone(tmp_path):
    source = build_source(
        tmp_path / "in" / "lista.xlsx", rows=[["P500", "Barato", "X", Decimal("0.25"), "$"]]
    )

    result = run(source, tmp_path, drift_share=Decimal("1"), anomaly=True, min_price=Decimal("1"))

    codes = {item["code"] for item in result.manifest["price_changes"]}
    assert "P500" not in codes and result.manifest["anomaly"]["code"] != "P500"
    assert products(table(result.path))["P500"] == Decimal("0.25")
