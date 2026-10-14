import csv
import io
from decimal import Decimal

from openpyxl import load_workbook

from pricor.adapters.csv_exporter import CsvExporter
from pricor.adapters.xlsx_exporter import XlsxExporter
from pricor.domain.contract import ContractRecord
from pricor.ports.exporter import CONTRACT_COLUMNS, CONVERTED_COLUMNS
from pricor.ports.price_history import ImportedRecord

from .exports import BARE, sample

CONVERTED = ContractRecord(
    "E1",
    "Antenna",
    Decimal("2.50"),
    "USD",
    unit_cost_net_converted=Decimal("3837.50"),
    fx_rate=Decimal("1535"),
    fx_source="dolarapi:oficial:venta",
    fx_date="2026-10-14",
)
PESOS = ContractRecord("E2", "Cable", Decimal("100.00"), "ARS")


def view():
    return sample([ImportedRecord(2, CONVERTED), ImportedRecord(3, PESOS)])


def csv_rows(content: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(content.decode("utf-8-sig"), newline="")))


def test_the_converted_columns_are_named_in_the_port():
    assert CONVERTED_COLUMNS == ("unit_cost_net_converted", "fx_rate", "fx_source", "fx_date")


def test_csv_adds_the_converted_columns_between_the_contract_and_the_source_row():
    rows = csv_rows(CsvExporter().export(view()))

    assert rows[0] == [*CONTRACT_COLUMNS, *CONVERTED_COLUMNS, "source_row"]
    assert rows[1][2:4] == ["2.50", "USD"]  # the original cost and currency stay put
    assert rows[1][-5:] == ["3837.50", "1535", "dolarapi:oficial:venta", "2026-10-14", "2"]
    assert rows[2][-5:] == ["", "", "", "", "3"]  # rows already in pesos have them blank


def test_csv_without_any_conversion_keeps_the_original_layout():
    rows = csv_rows(CsvExporter().export(sample([ImportedRecord(8, BARE)])))

    assert rows[0] == [*CONTRACT_COLUMNS, "source_row"]


def test_xlsx_adds_numeric_converted_cells_and_a_text_source():
    sheet = load_workbook(io.BytesIO(XlsxExporter().export(view())))["Products"]

    header = [cell.value for cell in sheet[1]]
    assert header == [*CONTRACT_COLUMNS, *CONVERTED_COLUMNS, "source_row"]
    column = {name: header.index(name) + 1 for name in header}
    converted = sheet.cell(row=2, column=column["unit_cost_net_converted"])
    assert (converted.value, converted.data_type, converted.number_format) == (3837.5, "n", "0.00")
    assert sheet.cell(row=2, column=column["fx_source"]).value == "dolarapi:oficial:venta"
    assert sheet.cell(row=2, column=column["fx_source"]).data_type == "s"
    assert sheet.cell(row=3, column=column["unit_cost_net_converted"]).value is None


def test_xlsx_without_any_conversion_keeps_the_original_layout():
    sheet = load_workbook(io.BytesIO(XlsxExporter().export(sample())))["Products"]

    assert [cell.value for cell in sheet[1]] == [*CONTRACT_COLUMNS, "source_row"]
