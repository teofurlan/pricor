from datetime import date, datetime
from decimal import Decimal

import openpyxl
import pytest
import xlwt

from pricor.adapters.spreadsheet_reader import FileSpreadsheetReader
from pricor.domain.engine import execute_plan
from pricor.domain.plan import parse_plan
from pricor.ports.spreadsheet import (
    FileTooLargeError,
    SheetNotFoundError,
    SpreadsheetError,
    TooManyRowsError,
    UnsupportedFormatError,
)


def _assert_no_float(rows):
    for row in rows:
        for cell in row:
            assert not isinstance(cell, float), cell


def _write_xlsx(path, sheets):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    workbook.save(path)


# -- CSV --


def test_csv_semicolon_utf8_with_bom(tmp_path):
    path = tmp_path / "list.csv"
    path.write_bytes("﻿Código;Precio\nA1;12,50\n".encode("utf-8"))
    sheet = FileSpreadsheetReader().read_sheet(path)
    assert sheet.rows == (("Código", "Precio"), ("A1", "12,50"))


def test_csv_cp1252_fallback(tmp_path):
    path = tmp_path / "list.csv"
    path.write_bytes("Descripción;Precio\nCañería;10\n".encode("cp1252"))
    sheet = FileSpreadsheetReader().read_sheet(path)
    assert sheet.rows[1] == ("Cañería", "10")


@pytest.mark.parametrize("delimiter", [";", ",", "\t"])
def test_csv_delimiter_sniffing(tmp_path, delimiter):
    path = tmp_path / "list.csv"
    lines = ["code,name,price", "A1,Tornillo,1", "A2,Tuerca,2"]
    path.write_text("\n".join(line.replace(",", delimiter) for line in lines), encoding="utf-8")
    sheet = FileSpreadsheetReader().read_sheet(path)
    assert sheet.rows[0] == ("code", "name", "price")
    assert len(sheet.rows) == 3


def test_csv_decimal_comma_does_not_fool_semicolon_sniffing(tmp_path):
    path = tmp_path / "list.csv"
    path.write_text("code;price\nA1;1,5\nA2;2,5\nA3;3,5\n", encoding="utf-8")
    sheet = FileSpreadsheetReader().read_sheet(path)
    assert sheet.rows[1] == ("A1", "1,5")


def test_csv_blank_cells_become_none(tmp_path):
    path = tmp_path / "list.csv"
    path.write_text("a;b;c\n1;;  \n", encoding="utf-8")
    sheet = FileSpreadsheetReader().read_sheet(path)
    assert sheet.rows[1] == ("1", None, None)


def test_csv_sheet_names_is_single_entry(tmp_path):
    path = tmp_path / "list.csv"
    path.write_text("a;b\n1;2\n", encoding="utf-8")
    assert FileSpreadsheetReader().sheet_names(path) == ["list"]


# -- XLSX / XLSM --


def test_xlsx_float_cell_becomes_exact_decimal(tmp_path):
    path = tmp_path / "list.xlsx"
    _write_xlsx(path, {"S": [["price"], [0.1], [12.5], [3]]})
    rows = FileSpreadsheetReader().read_sheet(path).rows
    assert rows[1] == (Decimal("0.1"),)
    assert isinstance(rows[1][0], Decimal)
    assert rows[2] == (Decimal("12.5"),)
    assert rows[3] == (3,) and isinstance(rows[3][0], int)
    _assert_no_float(rows)


def test_xlsx_dates_text_and_blanks(tmp_path):
    path = tmp_path / "list.xlsx"
    _write_xlsx(path, {"S": [["when", "gap", "name"], [datetime(2026, 10, 1), None, "Ana"]]})
    rows = FileSpreadsheetReader().read_sheet(path).rows
    assert rows[1][0] in (datetime(2026, 10, 1), date(2026, 10, 1))
    assert rows[1][1:] == (None, "Ana")


def test_xlsx_booleans_are_text_not_ints(tmp_path):
    path = tmp_path / "list.xlsx"
    _write_xlsx(path, {"S": [[True]]})
    assert FileSpreadsheetReader().read_sheet(path).rows == (("TRUE",),)


def test_xlsx_default_is_first_non_empty_sheet(tmp_path):
    path = tmp_path / "list.xlsx"
    _write_xlsx(path, {"Empty": [], "Prices": [["a"], ["b"]], "Other": [["z"]]})
    reader = FileSpreadsheetReader()
    sheet = reader.read_sheet(path)
    assert sheet.name == "Prices"
    assert reader.sheet_names(path) == ["Empty", "Prices", "Other"]


def test_xlsx_choose_sheet_by_name(tmp_path):
    path = tmp_path / "list.xlsx"
    _write_xlsx(path, {"A": [["a"]], "B": [["b"]]})
    sheet = FileSpreadsheetReader().read_sheet(path, "B")
    assert sheet.name == "B" and sheet.rows == (("b",),)


def test_unknown_sheet_name_is_an_error(tmp_path):
    path = tmp_path / "list.xlsx"
    _write_xlsx(path, {"A": [["a"]]})
    with pytest.raises(SheetNotFoundError):
        FileSpreadsheetReader().read_sheet(path, "Nope")


def test_workbook_without_content_is_an_error(tmp_path):
    path = tmp_path / "list.xlsx"
    _write_xlsx(path, {"A": []})
    with pytest.raises(SpreadsheetError, match="no non-empty sheet"):
        FileSpreadsheetReader().read_sheet(path)


def test_xlsm_extension_is_read_like_xlsx(tmp_path):
    source = tmp_path / "list.xlsx"
    _write_xlsx(source, {"S": [["a"], [1]]})
    path = tmp_path / "list.xlsm"
    path.write_bytes(source.read_bytes())
    assert FileSpreadsheetReader().read_sheet(path).rows == (("a",), (1,))


def test_corrupt_xlsx_raises_clear_error(tmp_path):
    path = tmp_path / "list.xlsx"
    path.write_bytes(b"not a zip file")
    with pytest.raises(SpreadsheetError, match="cannot read"):
        FileSpreadsheetReader().read_sheet(path)


# -- XLS --


def _write_xls(path):
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Old")
    sheet.write(0, 0, "price")
    sheet.write(1, 0, 0.1)
    sheet.write(2, 0, 7791234567890.0)
    sheet.write(3, 0, "text")
    workbook.save(str(path))


def test_xls_numbers_are_decimal_without_float_artifacts(tmp_path):
    path = tmp_path / "list.xls"
    _write_xls(path)
    rows = FileSpreadsheetReader().read_sheet(path).rows
    assert rows[1] == (Decimal("0.1"),)
    assert rows[2] == (7791234567890,) and isinstance(rows[2][0], int)
    assert rows[3] == ("text",)
    _assert_no_float(rows)


def test_content_decides_the_format_when_extension_lies(tmp_path):
    """Real supplier files exist with an .xlsx name but an OLE2 (.xls) body."""
    path = tmp_path / "real.xls"
    _write_xls(path)
    misnamed = tmp_path / "misnamed.xlsx"
    misnamed.write_bytes(path.read_bytes())
    assert FileSpreadsheetReader().read_sheet(misnamed).rows[1] == (Decimal("0.1"),)


def test_encrypted_workbook_gets_a_specific_error(tmp_path):
    """Password-protected xlsx files are OLE2 containers holding an EncryptedPackage stream."""
    path = tmp_path / "locked.xlsx"
    path.write_bytes(
        bytes.fromhex("d0cf11e0a1b11ae1") + "EncryptedPackage".encode("utf-16le") + bytes(64)
    )
    with pytest.raises(SpreadsheetError, match="password-protected"):
        FileSpreadsheetReader().read_sheet(path)


def test_xls_reader_does_not_write_to_stdout(tmp_path, capsys):
    path = tmp_path / "list.xls"
    _write_xls(path)
    FileSpreadsheetReader().read_sheet(path)
    assert capsys.readouterr().out == ""


# -- percentage semantics through the public reader and engine --


def _write_formatted_cells(path, values, formats):
    if path.suffix == ".xls":
        workbook = xlwt.Workbook()
        sheet = workbook.add_sheet("Prices")
        for index, (value, number_format) in enumerate(zip(values, formats, strict=True)):
            sheet.write(0, index, value, xlwt.easyxf(num_format_str=number_format))
        workbook.save(str(path))
    else:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for index, (value, number_format) in enumerate(zip(values, formats, strict=True), 1):
            cell = sheet.cell(1, index, value)
            cell.number_format = number_format
        workbook.save(path)


@pytest.mark.parametrize("extension", [".xlsx", ".xlsm", ".xls"])
@pytest.mark.parametrize(
    "price,discount,vat,expected",
    [(121, 0, 0.21, "100.00"), (100, 0.21, 0, "79.00")],
    ids=["vat", "discount"],
)
def test_native_excel_percentages_produce_correct_costs(
    tmp_path, capsys, extension, price, discount, vat, expected
):
    path = tmp_path / f"percentages{extension}"
    _write_formatted_cells(
        path,
        ["00123", "Widget", price, discount, vat],
        ["General", "General", "General", "0%", "0.00%"],
    )
    sheet = FileSpreadsheetReader().read_sheet(path)
    assert sheet.rows[0][3:] == (
        f"{Decimal(str(discount)) * 100:f}%",
        f"{Decimal(str(vat)) * 100:f}%",
    )
    _assert_no_float(sheet.rows)
    headers = ["Code", "Description", "Price", "Discount", "VAT"]
    plan = parse_plan({
        "contract_version": "1",
        "engine_version": "0.1.0",
        "supplier": "Percentage fixture",
        "steps": [
            {"op": "normalize_text", "column": {"index": 0, "header": "Code"},
             "to": "supplier_code"},
            {"op": "normalize_text", "column": {"index": 1, "header": "Description"},
             "to": "description"},
            {"op": "parse_number", "column": {"index": 2, "header": "Price"},
             "decimal_sep": ".", "thousands_sep": ""},
            {"op": "apply_discounts", "from": {
                "column": {"index": 3, "header": "Discount"}, "format": "plus_separated"}},
            {"op": "strip_vat", "rate_from": {"column": {"index": 4, "header": "VAT"}}},
            {"op": "constant", "field": "currency", "value": "ARS"},
        ],
        "rounding": {"decimals": 2, "mode": "half_up"},
        "confirmed_by_user": ["VAT and discount columns confirmed"],
    })
    result = execute_plan(plan, headers, sheet.rows)
    assert result.rejected == ()
    assert len(result.accepted) == 1
    record = result.accepted[0].record
    assert (record.supplier_code, record.description, record.currency) == ("00123", "Widget", "ARS")
    assert record.unit_cost_net == Decimal(expected)
    assert capsys.readouterr() == ("", "")


@pytest.mark.parametrize("extension", [".xlsx", ".xlsm", ".xls"])
def test_percent_formats_leave_unformatted_decimals_and_text_unchanged(tmp_path, extension):
    path = tmp_path / f"mixed{extension}"
    values = [0.21, "21%", 0.21, 0.21]
    _write_formatted_cells(path, values, ["General", "0%", '0"%"', r"0\%"])
    rows = FileSpreadsheetReader().read_sheet(path).rows
    assert rows == ((Decimal("0.21"), "21%", Decimal("0.21"), Decimal("0.21")),)
    _assert_no_float(rows)


@pytest.mark.parametrize("extension", [".xlsx", ".xls"])
@pytest.mark.parametrize(
    "value,number_format,expected",
    [
        (0.21, '"rate;"0.0%', "21.00%"),
        (0.21, "0%;0.00", "21.00%"),
        (-0.21, "0%;0.00", Decimal("-0.21")),
        (0.21, "0.00;0%", Decimal("0.21")),
        (-0.21, "0.00;0%", "-21.00%"),
        (0, "0%;0%;0", 0),
        (0.21, "[>=1]0%;0.00", Decimal("0.21")),
        (2, "[>=1]0%;0.00", "200%"),
        (0.105, "[Red]0%", "10.500%"),
    ],
)
def test_only_the_active_numeric_format_section_sets_percentage_semantics(
    tmp_path, extension, value, number_format, expected
):
    path = tmp_path / f"sections{extension}"
    _write_formatted_cells(path, [value], [number_format])
    assert FileSpreadsheetReader().read_sheet(path).rows == ((expected,),)


def test_csv_text_percentages_remain_unchanged(tmp_path):
    path = tmp_path / "percentages.csv"
    path.write_text("price;vat;discount\n0.21;21%;10%+5%\n", encoding="utf-8")
    rows = FileSpreadsheetReader().read_sheet(path).rows
    assert rows == (("price", "vat", "discount"), ("0.21", "21%", "10%+5%"))


# -- limits and safety --


def test_unsupported_extension(tmp_path):
    path = tmp_path / "list.pdf"
    path.write_bytes(b"%PDF")
    with pytest.raises(UnsupportedFormatError, match=".pdf"):
        FileSpreadsheetReader().read_sheet(path)


def test_file_size_limit(tmp_path):
    path = tmp_path / "list.csv"
    path.write_text("a;b\n1;2\n", encoding="utf-8")
    with pytest.raises(FileTooLargeError):
        FileSpreadsheetReader(max_file_bytes=4).read_sheet(path)


def test_row_limit_csv(tmp_path):
    path = tmp_path / "list.csv"
    path.write_text("a;b\n" + "1;2\n" * 10, encoding="utf-8")
    with pytest.raises(TooManyRowsError):
        FileSpreadsheetReader(max_rows=5).read_sheet(path)


def test_row_limit_xlsx(tmp_path):
    path = tmp_path / "list.xlsx"
    _write_xlsx(path, {"S": [[i] for i in range(10)]})
    with pytest.raises(TooManyRowsError):
        FileSpreadsheetReader(max_rows=5).read_sheet(path)


def test_trailing_empty_rows_are_trimmed_and_not_counted(tmp_path):
    path = tmp_path / "list.csv"
    path.write_text("a;b\n1;2\n" + ";\n" * 4, encoding="utf-8")
    sheet = FileSpreadsheetReader(max_rows=3).read_sheet(path)
    assert len(sheet.rows) == 2


def test_formula_like_text_is_kept_as_inert_text(tmp_path):
    path = tmp_path / "list.csv"
    path.write_text("a;b\n=1+1;@SUM(A1)\n", encoding="utf-8")
    assert FileSpreadsheetReader().read_sheet(path).rows[1] == ("=1+1", "@SUM(A1)")
