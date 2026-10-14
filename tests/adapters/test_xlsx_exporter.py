import io
import zipfile
from decimal import Decimal

from openpyxl import load_workbook

from pricor.adapters.xlsx_exporter import XlsxExporter
from pricor.domain.contract import ContractRecord
from pricor.ports.exporter import PRODUCT_COLUMNS, RejectedEntry
from pricor.ports.price_history import ImportedRecord

from .exports import sample


def workbook(content: bytes):
    return load_workbook(io.BytesIO(content))


def sheet_xml(content: bytes, number: int = 1) -> str:
    return zipfile.ZipFile(io.BytesIO(content)).read(f"xl/worksheets/sheet{number}.xml").decode()


def test_export_has_the_products_sheet_and_a_rejected_rows_sheet():
    book = workbook(XlsxExporter().export(sample()))
    assert book.sheetnames == ["Products", "Rejected rows"]
    assert [cell.value for cell in book["Products"][1]] == list(PRODUCT_COLUMNS)


def test_codes_and_barcodes_are_text_so_leading_zeros_survive():
    sheet = workbook(XlsxExporter().export(sample()))["Products"]
    code, barcode = sheet["A2"], sheet["G2"]
    assert (code.value, code.data_type, code.number_format) == ("00123", "s", "@")
    assert (barcode.value, barcode.data_type) == ("0779123456789", "s")


def test_numbers_are_numeric_cells_with_exact_decimal_values():
    sheet = workbook(XlsxExporter().export(sample()))["Products"]
    cost, vat, units, retail, source_row = (sheet[ref] for ref in ("C2", "E2", "F2", "H2", "I2"))
    for cell in (cost, vat, units, retail, source_row):
        assert cell.data_type == "n"
    assert Decimal(str(cost.value)) == Decimal("855.00")
    assert Decimal(str(vat.value)) == Decimal("10.5")
    assert Decimal(str(retail.value)) == Decimal("1234.5")
    assert cost.number_format == "0.00"  # 855.00 is displayed with its two decimals
    assert (units.value, source_row.value) == (12, 3)
    assert Decimal(str(sheet["C3"].value)) == Decimal("0.1")  # no 0.1000000000000000055


def test_numbers_are_written_with_their_own_digits_in_the_file():
    content = XlsxExporter().export(sample())
    xml = sheet_xml(content)
    assert "<v>0.1</v>" in xml
    assert "<v>1234.5</v>" in xml


def test_values_a_spreadsheet_number_cannot_hold_exactly_are_written_as_text():
    precise = ContractRecord("A", "B", Decimal("12345678901234567.89"), "ARS")
    cell = workbook(XlsxExporter().export(sample([ImportedRecord(1, precise)])))["Products"]["C2"]
    assert (cell.data_type, cell.value) == ("s", "12345678901234567.89")


def test_missing_optional_values_are_empty_cells():
    sheet = workbook(XlsxExporter().export(sample()))["Products"]
    assert [sheet.cell(row=3, column=n).value for n in (5, 6, 7, 8)] == [None] * 4


def test_untrusted_text_is_never_stored_as_a_formula():
    record = ContractRecord("=1+1", "=HYPERLINK(\"http://evil\")", Decimal("1"), "ARS", barcode="@x")
    rejected = RejectedEntry(2, ("=cmd", "+1"), ("=reason",))
    content = XlsxExporter().export(sample([ImportedRecord(1, record)], [rejected]))
    book = workbook(content)
    for sheet in book:
        assert [c.value for row in sheet.iter_rows() for c in row if c.data_type == "f"] == []
    assert book["Products"]["B2"].value == '=HYPERLINK("http://evil")'  # kept verbatim, as text
    assert book["Rejected rows"]["C2"].value == "=cmd"


def test_rejected_sheet_lists_every_rejected_row():
    data = sample()
    sheet = workbook(XlsxExporter().export(data))["Rejected rows"]
    rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert rows[0] == ["source_row", "reasons", "Codigo", "Descripcion", "Precio"]
    assert rows[1] == [5, "Precio: not a number", "A9", "Brida", "abc"]
    assert rows[2] == [6, "supplier_code is empty; x", None, "Sin codigo", "10,00"]
    assert len(rows) - 1 == len(data.rejected)


def test_standalone_rejected_report_is_one_sheet():
    book = workbook(XlsxExporter().export_rejected(sample()))
    assert book.sheetnames == ["Rejected rows"]
    assert book["Rejected rows"].max_row == 3


def test_original_cells_stay_text_in_the_report():
    data = sample(rejected=[RejectedEntry(4, ("00045", "10,50"), ("bad",))], headers=("Cod", "Precio"))
    sheet = workbook(XlsxExporter().export_rejected(data))["Rejected rows"]
    assert (sheet["C2"].value, sheet["C2"].data_type) == ("00045", "s")
    assert sheet["D2"].value == "10,50"


def test_no_rows_still_writes_headers():
    book = workbook(XlsxExporter().export(sample(accepted=[], rejected=[])))
    assert book["Products"].max_row == 1
    assert book["Rejected rows"].max_row == 1


def test_metadata():
    exporter = XlsxExporter()
    assert (exporter.format, exporter.extension) == ("xlsx", ".xlsx")
    assert "spreadsheetml" in exporter.media_type
